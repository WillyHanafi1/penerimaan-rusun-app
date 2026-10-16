import unittest
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd
import openpyxl

import app
from kasda_matching import read_kasda_sewa, reconcile_kasda_sewa, _period, _unit_code, _va_reference


def bank_row(ref='ABC123', amount=400000, **kwargs):
    return {'Narasi': f'ESB-V001-{ref} Payment QRIS RETRIBUSI',
            'Posting Date': '04 Sep 2026', 'Credit Transaction': amount,
            'Debit Transaction': '', **kwargs}


def receipt(ref='000000ABC123', amount=400000, time='08:00:00', **kwargs):
    return {'Baris_Kasda': 9, 'Kode_8_Digit': '01010101', 'Bulan': 'September',
            'Tahun': '2026', 'Tanggal_Kasda': pd.Timestamp('2026-09-04'),
            'Waktu_Bayar_Kasda': time, 'Metode_Bayar_Kasda': 'QRIS',
            'No_Referensi_Kasda': ref, 'Nama_Kasda': 'Contoh',
            'Jumlah_Kasda': amount, 'Denda_Kasda': 0,
            'Keterangan_Pencocokan': '', **kwargs}


def match(bank_rows, receipts):
    detected, manual = app.process_setortunai(pd.DataFrame(bank_rows))
    return reconcile_kasda_sewa(detected, manual, pd.DataFrame(receipts))


class KasdaMatchingTests(unittest.TestCase):
    def daily_workbook(self, names=('13-10-2026', '14-10-2026', '15-10-2026'), empty=False):
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for name in names:
            sheet = workbook.create_sheet(name)
            sheet.append(['RUSUNAWA CIBEUREUM TYPE 24'])
            sheet.append(['NO. URUT', 'Tanggal Setor', 'Tanggal Kasda', 'NOMOR HUNIAN', None, None,
                          'NAMA PENGHUNI', 'TANGGAL PERJANJIAN', 'SEWA HUNIAN', 'PARKIR', 'DENDA',
                          'JUMLAH', 'BULAN', 'KANAL BAYAR', 'WAKTU BAYAR (WIB)', 'NO. REFERENSI',
                          'NOMOR SKRD / STRD', 'BIAYA LAIN'])
            sheet.merge_cells('D2:F2')
            if not empty:
                day = int(name.split('-')[0])
                sheet.append([1, datetime(2026, 10, day), datetime(2026, 10, day + 1),
                              'D', 'V', 24, 'Contoh', datetime(2025, 1, 1), 350000, 30000,
                              15000, 400000, datetime(2025, 12, 1), 'QRIS',
                              datetime(2026, 10, day, 8, 30), '1295012345678901', 'SKRD/1', 5000])
                sheet.append(['TOTAL', None, None, None, None, None, None, None, None, None, None, 400000])
        workbook.create_sheet('Denda').append(['Do not read this sheet'])
        return workbook

    def test_daily_sheet_is_selected_by_bank_date_regardless_of_position(self):
        for names in (('13-10-2026', '14-10-2026', '15-10-2026'),
                      ('14-10-2026', '15-10-2026', '13-10-2026'),
                      ('15-10-2026', '13-10-2026', '14-10-2026')):
            with patch('kasda_matching.openpyxl.load_workbook', return_value=self.daily_workbook(names)):
                receipts = read_kasda_sewa('daily.xlsx', ['14 Oct 2026'])
            self.assertEqual(len(receipts), 1)
            row = receipts.iloc[0]
            self.assertEqual(row['Sheet_Kasda'], '14-10-2026')
            self.assertEqual(row['Kode_8_Digit'], '02040524')
            self.assertEqual((row['Bulan'], row['Tahun']), ('Desember', '2025'))
            self.assertEqual(row['Waktu_Bayar_Kasda'], '08:30:00')
            self.assertEqual(row['No_Referensi_Kasda'], '1295012345678901')
            self.assertEqual(row['Biaya_Lain_Kasda'], 5000)
            self.assertEqual(row['Nomor_Dokumen_Kasda'], 'SKRD/1')

    def test_daily_matching_uses_payment_date_and_retains_separate_kasda_date(self):
        with patch('kasda_matching.openpyxl.load_workbook', return_value=self.daily_workbook()):
            receipts = read_kasda_sewa('daily.xlsx', ['14 Oct 2026'])
        bank = pd.DataFrame([bank_row('678901', **{'Posting Date': '14 Oct 2026'})])
        detected, manual = app.process_setortunai(bank)
        detected, manual, audit = reconcile_kasda_sewa(detected, manual, receipts)
        self.assertEqual(len(detected), 1)
        self.assertTrue(manual.empty)
        self.assertEqual(detected.iloc[0]['Tanggal_Kasda'], pd.Timestamp('2026-10-15'))
        self.assertEqual(detected.iloc[0]['Tanggal_Bayar_Kasda'], pd.Timestamp('2026-10-14'))
        self.assertEqual(audit.iloc[0]['Sheet_Kasda'], '14-10-2026')

    def test_daily_multiple_bank_dates_select_only_the_requested_sheets(self):
        with patch('kasda_matching.openpyxl.load_workbook', return_value=self.daily_workbook()):
            receipts = read_kasda_sewa('daily.xlsx', ['13 Oct 2026', '15 Oct 2026'])
        self.assertEqual(receipts['Sheet_Kasda'].tolist(), ['13-10-2026', '15-10-2026'])

    def test_daily_missing_date_or_invalid_header_is_reported_instead_of_using_last_sheet(self):
        with patch('kasda_matching.openpyxl.load_workbook', return_value=self.daily_workbook()):
            with self.assertRaisesRegex(ValueError, 'tidak tersedia'):
                read_kasda_sewa('daily.xlsx', ['12 Oct 2026'])
        workbook = self.daily_workbook()
        workbook['14-10-2026']['N2'] = 'Invalid payment header'
        with patch('kasda_matching.openpyxl.load_workbook', return_value=workbook):
            with self.assertRaisesRegex(ValueError, 'Header penerimaan'):
                read_kasda_sewa('daily.xlsx', ['14 Oct 2026'])

    def test_valid_empty_daily_sheet_keeps_bank_transactions_manual(self):
        with patch('kasda_matching.openpyxl.load_workbook', return_value=self.daily_workbook(empty=True)):
            receipts = read_kasda_sewa('daily.xlsx', ['14 Oct 2026'])
        detected, manual = app.process_setortunai(pd.DataFrame([bank_row(**{'Posting Date': '14 Oct 2026'})]))
        detected, manual, audit = reconcile_kasda_sewa(detected, manual, receipts)
        self.assertTrue(detected.empty)
        self.assertEqual(len(manual), 1)
        self.assertTrue(audit.empty)

    def test_reference_uses_amount_and_date_and_retains_bank_values(self):
        detected, manual, audit = match([bank_row()], [receipt()])
        self.assertEqual(len(detected), 1)
        self.assertTrue(manual.empty)
        self.assertEqual(detected.iloc[0]['Credit Transaction'], 400000)
        self.assertEqual(detected.iloc[0]['Narasi'], bank_row()['Narasi'])
        self.assertEqual(audit.iloc[0]['Status_Pencocokan'], 'Cocok')
        for other in [receipt(amount=390000), receipt(Tanggal_Kasda=pd.Timestamp('2026-09-05'))]:
            self.assertTrue(match([bank_row()], [other])[0].empty)

    def test_time_order_fallback_is_ascending(self):
        rows = [receipt(ref='WRONG2', time='09:00:00', Kode_8_Digit='01010102'),
                receipt(ref='WRONG1', time='08:00:00')]
        detected, _, audit = match([bank_row('FIRST1'), bank_row('LAST01')], rows)
        self.assertEqual(detected['Kode_8_Digit'].tolist(), ['01010101', '01010102'])
        self.assertEqual(audit['ID_Transaksi'].nunique(), 2)

    def test_missing_receipt_cannot_shift_rows_or_reuse_a_receipt(self):
        detected, manual, _ = match([bank_row('FIRST1'), bank_row()], [receipt()])
        self.assertEqual(detected['ID_Transaksi'].tolist(), ['BANK-000002'])
        self.assertEqual(len(manual), 1)

    def test_debit_does_not_consume_time_order_position(self):
        detected, manual, _ = match([
            bank_row('FIRST1'), bank_row(**{'Credit Transaction': '', 'Debit Transaction': 400000}),
            bank_row('LAST01')], [receipt(ref='WRONG1'), receipt(ref='WRONG2', time='09:00:00')])
        self.assertEqual(detected['ID_Transaksi'].tolist(), ['BANK-000001', 'BANK-000003'])
        self.assertEqual(len(manual), 1)

    def test_conflicting_reference_anchors_disable_fallback(self):
        detected, manual, _ = match([bank_row(), bank_row('LAST01')],
            [receipt(ref='OTHER1', time='08:00:00'), receipt(time='09:00:00')])
        self.assertEqual(len(detected), 1)
        self.assertEqual((manual['Jenis_Baris'] == 'Transaksi bank').sum(), 1)
        self.assertEqual((manual['Jenis_Baris'] == 'Kasda tanpa pasangan bank').sum(), 1)
        self.assertEqual(detected.iloc[0]['Waktu_Bayar_Kasda'], '09:00:00')

    def test_duplicate_references_and_tied_times_are_not_guessed(self):
        detected, _, audit = match([bank_row(), bank_row()], [receipt(), receipt()])
        self.assertTrue(detected.empty)
        self.assertTrue((audit['Status_Pencocokan'] == 'Belum cocok').all())

    def test_partial_reference_matches_work_when_recap_has_extra_rows(self):
        detected, _, audit = match([bank_row()], [receipt(), receipt(ref='EXTRA1')])
        self.assertEqual(len(detected), 1)
        self.assertEqual((audit['Status_Pencocokan'] == 'Belum cocok').sum(), 1)

    def test_invalid_receipt_goes_manual_without_blocking_valid_rows(self):
        detected, manual, _ = match([bank_row('OTHER1')], [receipt(), receipt(
            Tanggal_Kasda=pd.NaT, Keterangan_Pencocokan='Tanggal Kasda tidak valid')])
        self.assertEqual(len(detected), 1)
        self.assertEqual(len(manual), 1)
        self.assertEqual(manual.iloc[0]['Jenis_Baris'], 'Kasda tanpa pasangan bank')

    def test_amount_mismatch_skips_bank_row_and_continues(self):
        detected, manual, _ = match([
            bank_row('FIRST1', 400000), bank_row('EXTRA1', 123456), bank_row('LAST01', 410000)],
            [receipt(ref='WRONG1', amount=400000), receipt(ref='WRONG2', amount=410000, time='09:00:00')])
        self.assertEqual(detected['ID_Transaksi'].tolist(), ['BANK-000001', 'BANK-000003'])
        self.assertEqual(manual['ID_Transaksi'].tolist(), ['BANK-000002'])
        self.assertEqual(detected['Credit Transaction'].sum(), 810000)

    def test_amount_mismatch_can_skip_an_extra_kasda_row(self):
        detected, manual, _ = match([bank_row('FIRST1', 400000), bank_row('LAST01', 410000)], [
            receipt(ref='WRONG1', amount=400000),
            receipt(ref='EXTRA1', amount=123456, time='08:30:00', Baris_Kasda=99),
            receipt(ref='WRONG2', amount=410000, time='09:00:00')])
        self.assertEqual(len(detected), 2)
        self.assertEqual(len(manual), 1)
        self.assertEqual(manual.iloc[0]['Jenis_Baris'], 'Kasda tanpa pasangan bank')
        self.assertTrue(pd.isna(manual.iloc[0]['Credit Transaction']))
        self.assertEqual(manual.iloc[0]['Jumlah_Kasda'], 123456)

    def test_ambiguous_repeated_amount_does_not_block_later_unique_match(self):
        detected, manual, _ = match([
            bank_row('FIRST1', 400000), bank_row('OTHER1', 400000), bank_row('LAST01', 410000)],
            [receipt(ref='WRONG1', amount=400000), receipt(ref='WRONG2', amount=410000, time='09:00:00')])
        self.assertEqual(detected['ID_Transaksi'].tolist(), ['BANK-000003'])
        self.assertEqual((manual['Jenis_Baris'] == 'Transaksi bank').sum(), 2)

    def test_crossing_amount_alternatives_are_not_both_accepted(self):
        detected, _, _ = match([bank_row('FIRST1', 400000), bank_row('LAST01', 410000)],
            [receipt(ref='WRONG1', amount=410000), receipt(ref='WRONG2', amount=400000, time='09:00:00')])
        self.assertTrue(detected.empty)

    def test_reference_amount_conflict_is_not_reassigned_to_another_receipt(self):
        detected, manual, _ = match([bank_row(amount=390000)],
            [receipt(amount=400000), receipt(ref='OTHER1', amount=390000, time='09:00:00')])
        self.assertTrue(detected.empty)
        self.assertIn('nominal bank dan Kasda berbeda', manual.iloc[0]['Keterangan'])

    def test_qris_settlement_is_not_a_payment_candidate(self):
        detected, manual, _ = match([bank_row(Narasi='D601/PELIMP VIA QRIS RETRIBUSI PENYEDOTAN KAKUS')], [receipt()])
        self.assertTrue(detected.empty)
        self.assertEqual((manual['Jenis_Baris'] == 'Transaksi bank').sum(), 1)

    def test_mobile_banking_va_is_matched_by_its_full_reference(self):
        detected, _, audit = match([bank_row(Narasi='MBL-M001-983616 PAY VA MBL 1295010202100826/Retribusi')],
            [receipt(ref='1295010202100826', Metode_Bayar_Kasda='VA')])
        self.assertEqual(len(detected), 1)
        self.assertTrue(audit.iloc[0]['Metode_Pencocokan'].startswith('Referensi'))

    def test_va_reference_and_previous_year_period(self):
        detected, _, _ = match([bank_row(Narasi='TLR-J580-412543 PAY VA TLR 1295030102171225/Retribusi')],
            [receipt(ref='1295030102171225', Metode_Bayar_Kasda='VA', Bulan='Desember', Tahun='2025')])
        self.assertEqual(detected.iloc[0]['Tahun'], '2025')
        self.assertEqual(detected.iloc[0]['Bulan'], 'Desember')

    def test_transfer_with_embedded_va_reference_is_matched(self):
        narration = 'ATM-JALN-303058 TRF DARI BCA Erni Murningsih NO REK 1295020104170726'
        detected, manual, audit = match(
            [bank_row(Narasi=narration, amount=390150)],
            [receipt(ref='1295020104170726', amount=390150, Metode_Bayar_Kasda='VA',
                     Kode_8_Digit='02010417', Bulan='Juli', Tahun='2026')],
        )
        self.assertEqual(len(detected), 1)
        self.assertTrue(manual.empty)
        self.assertEqual(detected.iloc[0]['Kode_8_Digit'], '02010417')
        self.assertEqual(detected.iloc[0]['Bulan'], 'Juli')
        self.assertEqual(detected.iloc[0]['Tahun'], '2026')
        self.assertTrue(audit.iloc[0]['Metode_Pencocokan'].startswith('Referensi'))

    def test_va_reference_requires_valid_unit_prefix_and_month(self):
        self.assertEqual(_va_reference('NO REK 1295020104170726'), '1295020104170726')
        self.assertEqual(_va_reference('NO REK 1295999999990726'), '')
        self.assertEqual(_va_reference('NO REK 1295020104171326'), '')

    def test_period_and_code_mapping(self):
        self.assertEqual(_period('Agu 26'), ('Agustus', '2026'))
        self.assertEqual(_period('Des 2025'), ('Desember', '2025'))
        self.assertEqual(_unit_code('RUSUNAWA CIBEUREUM TYPE 27', 'D', 'V', 24), '02040524')
        self.assertEqual(_unit_code('UNKNOWN', 'D', 'V', 24), '')

    def test_unsupported_year_and_high_credit_stay_manual(self):
        detected, manual, _ = match([bank_row(amount=750000)], [receipt(amount=750000, Tahun='2024')])
        eligible, manual = app.filter_incomplete_data(detected, manual, ['2025', '2026'])
        self.assertTrue(eligible.empty)
        self.assertIn('2024', manual.iloc[0]['Keterangan'])
        self.assertIn('700.000', manual.iloc[0]['Keterangan'])

    def test_without_kasda_preserves_existing_narration_path(self):
        bank = pd.DataFrame([bank_row(), bank_row(Narasi='SETOR TUNAI 01010101 DESEMBER2025')])
        detected, manual, audit = app.process_bank_sources(bank)
        self.assertEqual(len(detected), 1)
        self.assertEqual(len(manual), 1)
        self.assertEqual(detected.iloc[0]['Tahun'], '2025')
        self.assertIsNone(audit)

    def test_report_status_is_specific_to_each_bank_transaction(self):
        data = pd.DataFrame([{**receipt(), 'ID_Transaksi': f'BANK-{i}', 'Tahun': year,
                              'Posting Date': '04 Sep 2026', 'Credit Transaction': 400000,
                              'Sewa Hunian': 400000, 'Sewa Lahan Lantai 1': 0,
                              'Nama Penghuni': 'Contoh', 'Tanggal Perjanjian Sewa': ''}
                             for i, year in enumerate(['2025', '2026', '2026'])])
        final = app.calculate_denda(data)
        base = {'Kode_8_Digit': '01010101', 'Bulan': 'September'}
        results = {'success_details': [{**base, 'ID_Transaksi': 'BANK-0', 'Sheet': 'CIGUGUR', 'Row': 3, 'Denda': 0},
                                       {**base, 'ID_Transaksi': 'BANK-1', 'Sheet': 'CIGUGUR', 'Row': 3, 'Denda': 0}],
                   'skipped_details': [{**base, 'ID_Transaksi': 'BANK-2', 'Reason': 'Sudah terisi'}]}
        captured = {}
        def capture(frame, writer, sheet_name, **kwargs):
            captured[sheet_name] = frame.copy()
        with patch('pandas.ExcelWriter') as writer, patch.object(pd.DataFrame, 'to_excel', capture), patch('os.makedirs'):
            writer.return_value.__enter__.return_value.book = SimpleNamespace(sheetnames=[])
            app.create_export_excel(results, final, final, final.head(0))
        self.assertEqual(captured['Status_Input']['Status_Input'].tolist(),
                         ['Berhasil Input', 'Berhasil Input', 'Skip - Sudah Terisi'])

    def test_matched_receipts_write_to_their_year_and_skip_duplicate(self):
        workbooks = {}
        for year in ['2025', '2026']:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'CIGUGUR'
            for col in ['B', 'C', 'D']:
                ws[f'{col}3'] = '01'
            workbooks[year] = wb
        data = pd.DataFrame([{'ID_Transaksi': f'BANK-{i}', 'Kode_8_Digit': '01010101',
                              'Bulan': 'September', 'Tahun': year, 'Posting Date': '04 Sep 2026', 'Denda': denda}
                             for i, (year, denda) in enumerate([('2025', 5000), ('2026', 15000), ('2026', 25000)])])
        with patch('shutil.copy2'), patch('os.path.exists', return_value=True), \
             patch('openpyxl.load_workbook', side_effect=lambda p: workbooks['2025' if '2025' in p else '2026']), \
             patch('app.save_master') as save:
            _, result = app.input_to_excel_master(data, {'2025': 'master_2025.xlsx', '2026': 'master_2026.xlsx'})
        self.assertEqual((result['success'], result['skipped'], result['failed']), (2, 1, 0))
        self.assertEqual(save.call_count, 2)
        self.assertEqual(workbooks['2025']['CIGUGUR']['CO3'].value, 5000)
        self.assertEqual(workbooks['2026']['CIGUGUR']['CO3'].value, 15000)
        self.assertEqual(workbooks['2026']['CIGUGUR']['CR3'].value.strftime('%Y-%m-%d'), '2026-09-04')
        self.assertEqual(workbooks['2026']['CIGUGUR']['CT3'].value, 1)

    def test_missing_master_lookup_never_enters_status_or_master_input(self):
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.title = 'CIGUGUR'
        for column in ['B', 'C', 'D']:
            ws[f'{column}3'] = '01'
        ws['CK3'] = 'Contoh'
        ws['CM3'] = 395000
        data = pd.DataFrame([{'ID_Transaksi': f'BANK-{i}', 'Urutan_Rekening_Koran': i + 1,
                              'Kode_8_Digit': code, 'Bulan': 'September', 'Tahun': '2026',
                              'Posting Date': '04 Sep 2026', 'Credit Transaction': 400000}
                             for i, code in enumerate(['01010101', '01999999'])])
        with patch('openpyxl.load_workbook', return_value=workbook):
            extracted = app.extract_from_master_excel(data, {'2026': 'master_2026.xlsx'})
        eligible, manual = app.filter_unmatched_master(extracted, data.head(0))
        self.assertEqual(eligible['ID_Transaksi'].tolist(), ['BANK-0'])
        self.assertEqual(manual['ID_Transaksi'].tolist(), ['BANK-1'])
        final = app.calculate_denda(eligible)
        with patch('shutil.copy2'), patch('os.path.exists', return_value=True), \
             patch('openpyxl.load_workbook', return_value=workbook), \
             patch('app.save_master'):
            valid, results = app.input_to_excel_master(final, {'2026': 'master_2026.xlsx'})
        self.assertEqual(results['success'], 1)
        self.assertEqual(results['failed'], 0)
        captured = {}
        def capture(frame, writer, sheet_name, **kwargs):
            captured[sheet_name] = frame.copy()
        with patch('pandas.ExcelWriter') as writer, patch.object(pd.DataFrame, 'to_excel', capture), patch('os.makedirs'):
            writer.return_value.__enter__.return_value.book = SimpleNamespace(sheetnames=[])
            app.create_export_excel(results, valid, final, manual)
        self.assertEqual(captured['Status_Input']['ID_Transaksi'].tolist(), ['BANK-0'])
        self.assertEqual(captured['Cek Manual']['ID_Transaksi'].tolist(), ['BANK-1'])

    @unittest.skipUnless(Path('rekap-pendapatan-2026-09-04 (2).xlsx').exists() and Path('4-9-2026.pdf').exists(), 'Local example inputs unavailable')
    def test_local_recap_reconciles_every_credit_and_preserves_the_debit(self):
        kasda = read_kasda_sewa('rekap-pendapatan-2026-09-04 (2).xlsx')
        self.assertEqual(len(kasda), 49)  # Totals and Denda sheet are not receipts.
        bank = app.process_pdf('4-9-2026.pdf')
        detected, manual, audit = app.process_bank_sources(bank, 'rekap-pendapatan-2026-09-04 (2).xlsx')
        self.assertEqual(len(detected), 49)
        self.assertEqual(len(manual), 1)
        self.assertEqual(app.to_numeric_safe(manual.iloc[0]['Debit Transaction']), 18727370)
        self.assertEqual(detected['ID_Transaksi'].nunique(), 49)
        self.assertEqual(detected['Baris_Kasda'].nunique(), 49)
        self.assertEqual(detected['Credit Transaction'].apply(app.to_numeric_safe).sum(), 20027570)
        self.assertEqual((audit['Status_Pencocokan'] == 'Cocok').sum(), 49)
        self.assertEqual(audit['Metode_Pencocokan'].str.startswith('Urutan').sum(), 1)

    @unittest.skipUnless(
        Path('07-09-2026/rekap-pendapatan-2026-09-07.xlsx').exists()
        and Path('07-09-2026/7-9-2026.pdf').exists(),
        'Local September 7 inputs unavailable',
    )
    def test_september_7_recap_all_enters_eligible_status_flow(self):
        folder = Path('07-09-2026')
        bank = app.process_pdf(str(folder / '7-9-2026.pdf'))
        detected, manual, audit = app.process_bank_sources(
            bank, str(folder / 'rekap-pendapatan-2026-09-07.xlsx'))
        eligible, _ = app.filter_incomplete_data(detected, manual, ['2025', '2026'])

        matched_ids = set(audit.loc[audit['Status_Pencocokan'] == 'Cocok', 'ID_Transaksi'])
        self.assertEqual(len(audit), 52)
        self.assertEqual(len(matched_ids), 52)
        self.assertEqual((audit['Status_Pencocokan'] == 'Cocok').sum(), 52)
        self.assertTrue(matched_ids.issubset(set(eligible['ID_Transaksi'])))
        self.assertEqual(len(eligible), 53)  # 52 rekap + 1 setor tunai dengan kode lengkap.
        target = eligible[eligible['Kode_8_Digit'] == '02010417'].iloc[0]
        self.assertEqual((target['Bulan'], target['Tahun']), ('Juli', '2026'))

    @unittest.skipUnless(Path('1-9-2026/1-9-2026.pdf').exists(), 'Local September 1 inputs unavailable')
    def test_september_1_mixed_transactions_have_exclusive_destinations(self):
        folder = Path('1-9-2026')
        bank = app.process_pdf(str(folder / '1-9-2026.pdf'))
        detected, manual, audit = app.process_bank_sources(bank, str(folder / 'rekap-pendapatan-2026-09-01 (2).xlsx'))
        eligible, manual = app.filter_incomplete_data(detected, manual, ['2025', '2026'])
        self.assertEqual(len(eligible), 17)
        self.assertEqual(len(manual), 10)
        self.assertEqual((audit['Status_Pencocokan'] == 'Cocok').sum(), 14)
        self.assertEqual(set(manual['No']), {'3', '11', '12', '14', '19', '21', '22', '25', '26', '27'})
        self.assertFalse(set(eligible['ID_Transaksi']) & set(manual['ID_Transaksi']))
        self.assertEqual(eligible['Credit Transaction'].apply(app.to_numeric_safe).sum(), 6941300)

    @unittest.skipUnless(Path('1-9-2026/1-9-2026.pdf').exists(), 'Local September 1 inputs unavailable')
    def test_september_1_time_order_recovers_after_extra_rows_on_both_sides(self):
        folder = Path('1-9-2026')
        bank = app.process_pdf(str(folder / '1-9-2026.pdf'))
        extra = pd.DataFrame([bank_row('EXTRA1', 123456, **{'Posting Date': '01 Sep 2026', 'No': 'EXTRA'})])
        bank = pd.concat([bank.iloc[:8], extra, bank.iloc[8:]], ignore_index=True)
        kasda = read_kasda_sewa(str(folder / 'rekap-pendapatan-2026-09-01 (2).xlsx'))
        kasda['No_Referensi_Kasda'] = ''  # Force time-order matching on copies, never source files.
        extra_receipt = kasda.iloc[[0]].copy()
        extra_receipt['Baris_Kasda'] = 999
        extra_receipt['Waktu_Bayar_Kasda'] = '09:10:00'
        extra_receipt['Jumlah_Kasda'] = 654321
        kasda = pd.concat([kasda, extra_receipt], ignore_index=True)
        detected, manual = app.process_setortunai(bank)
        detected, manual, audit = reconcile_kasda_sewa(detected, manual, kasda)
        eligible, manual = app.filter_incomplete_data(detected, manual, ['2025', '2026'])
        self.assertEqual(len(eligible), 17)
        self.assertEqual((audit['Status_Pencocokan'] == 'Cocok').sum(), 14)
        self.assertTrue((audit.loc[audit['Status_Pencocokan'] == 'Cocok', 'Metode_Pencocokan'].str.startswith('Urutan')).all())
        self.assertEqual(len(manual), 12)
        self.assertIn('EXTRA', set(manual['No']))
        self.assertEqual((manual['Jenis_Baris'] == 'Kasda tanpa pasangan bank').sum(), 1)
        self.assertEqual(eligible['Credit Transaction'].apply(app.to_numeric_safe).sum(), 6941300)

    def test_master_input_payment_method_keterangan_and_fill_colors(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'CIGUGUR'
        for i, row in enumerate([3, 4, 5], 1):
            ws[f'B{row}'] = '01'
            ws[f'C{row}'] = '01'
            ws[f'D{row}'] = f'{i:02d}'
            ws[f'CK{row}'] = f'Penghuni {i}'
        
        ws['CQ5'] = 'Penghuni Baru'

        data = pd.DataFrame([
            {
                'ID_Transaksi': 'BANK-000001',
                'Kode_8_Digit': '01010101',
                'Bulan': 'September',
                'Tahun': '2026',
                'Posting Date': '01 Sep 2026',
                'Denda': 0,
                'Metode_Bayar_Kasda': 'QRIS',
            },
            {
                'ID_Transaksi': 'BANK-000002',
                'Kode_8_Digit': '01010102',
                'Bulan': 'September',
                'Tahun': '2026',
                'Posting Date': '01 Sep 2026',
                'Denda': 5000,
                'Metode_Bayar_Kasda': 'VA',
            },
            {
                'ID_Transaksi': 'BANK-000003',
                'Kode_8_Digit': '01010103',
                'Bulan': 'September',
                'Tahun': '2026',
                'Posting Date': '01 Sep 2026',
                'Denda': 0,
                'Metode_Bayar_Kasda': 'QRIS',
            }
        ])

        with patch('shutil.copy2'), patch('os.path.exists', return_value=True), \
             patch('openpyxl.load_workbook', return_value=wb), \
             patch.object(openpyxl.workbook.workbook.Workbook, 'save'):
            valid, results = app.input_to_excel_master(data, {'2026': 'master_2026.xlsx'})

        self.assertEqual(results['success'], 3)
        self.assertEqual(ws['CQ3'].value, 'QRIS')
        self.assertIn('C5D9F1', str(ws['CQ3'].fill.start_color.rgb))
        self.assertIn('C5D9F1', str(ws['CR3'].fill.start_color.rgb))

        self.assertEqual(ws['CQ4'].value, 'VA')
        self.assertIn('D9B5CA', str(ws['CQ4'].fill.start_color.rgb))
        self.assertIn('D9B5CA', str(ws['CR4'].fill.start_color.rgb))

        self.assertEqual(ws['CQ5'].value, 'Penghuni Baru, QRIS')
        self.assertIn('C5D9F1', str(ws['CQ5'].fill.start_color.rgb))


if __name__ == '__main__':
    unittest.main()
