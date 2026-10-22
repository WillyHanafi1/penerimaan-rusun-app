from io import BytesIO
from datetime import datetime
import hashlib
from pathlib import Path
import tempfile
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
import pandas as pd

import app
from master_formula_cache import NS, save_master


class MasterOutputTests(unittest.TestCase):
    def test_after_16_bank_matched_payment_reaches_master_and_replay_is_skipped(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / 'master.xlsx'
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = 'CIGUGUR'
            for column in ('B', 'C', 'D'):
                sheet[f'{column}3'] = '01'
            sheet['CU3'] = 'Penghuni Master'
            sheet['CW3'] = 300000
            workbook.save(source)
            original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            kasda = openpyxl.Workbook()
            ws = kasda.active
            ws.title = '14-10-2026'
            ws.append(['RUSUNAWA CIGUGUR TYPE 21'])
            ws.append(['NO. URUT', 'Tanggal Setor', 'Tanggal Kasda', 'GD', 'LT', 'UNIT',
                       'NAMA PENGHUNI', 'SEWA HUNIAN', 'PARKIR', 'DENDA', 'JUMLAH',
                       'BULAN', 'KANAL BAYAR', 'WAKTU BAYAR (WIB)', 'NO. REFERENSI'])
            ws.append([1, datetime(2026, 10, 14), datetime(2026, 10, 15),
                       'A', 'I', 1, 'Nama Kasda', 300000, 0, 10000, 310000,
                       'Oktober 2026', 'QRIS', '14-10-2026 23:59:59', '000000LATE01'])
            payload = BytesIO()
            kasda.save(payload)
            payload.seek(0)
            bank = pd.DataFrame([{'Narasi': 'ESB-V001-LATE01 Payment QRIS RETRIBUSI',
                                  'Posting Date': '14 Oct 2026', 'Credit Transaction': 310000,
                                  'Debit Transaction': ''}])
            detected, manual, audit = app.process_bank_sources(bank, payload)
            self.assertTrue(manual.empty)
            self.assertEqual(audit.iloc[0]['Waktu_Bayar_Kasda'], '23:59:59')
            eligible, manual = app.filter_incomplete_data(detected, manual, ['2026'])
            extracted = app.extract_from_master_excel(eligible, {'2026': str(source)})
            eligible, manual = app.filter_unmatched_master(extracted, manual)
            self.assertTrue(manual.empty)
            final = app.calculate_denda(eligible)
            valid, results = app.input_to_excel_master(final, {'2026': str(source)}, str(root / 'first'))
            self.assertEqual((results['success'], results['failed']), (1, 0))
            updated_path = Path(results['backup_files']['2026'])
            updated = openpyxl.load_workbook(updated_path)['CIGUGUR']
            self.assertEqual(updated['CU3'].value, 'Penghuni Master')
            self.assertEqual(updated['CY3'].value, 10000)
            self.assertEqual(updated['DA3'].value, 'QRIS')
            self.assertEqual(updated['DB3'].value, datetime(2026, 10, 14))
            report = app.create_export_excel(results, valid, final, manual, audit, str(root / 'first'))
            status = openpyxl.load_workbook(report)['Status_Input']
            headers = {c.value: c.column for c in status[1]}
            self.assertEqual(status.cell(2, headers['Status_Input']).value, 'Berhasil Input')
            _, replay = app.input_to_excel_master(final, {'2026': str(updated_path)}, str(root / 'second'))
            self.assertEqual((replay['success'], replay['skipped'], replay['failed']), (0, 1, 0))
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)

    def test_status_review_rule_keeps_only_plain_success_normal(self):
        cases = [
            ('Berhasil Input', 'Data berhasil diinput ke LG  row 273', False),
            ('Berhasil Input', 'Data   berhasil diinput ke MELONG row 222', False),
            ('Berhasil Input', 'Data berhasil diinput ke CIGUGUR row 35', False),
            ('Berhasil Input', 'Data berhasil diinput ke MELONG row 267; rincian sewa/denda memakai Kasda', True),
            ('Berhasil Input', 'Data berhasil diinput ke LG row 205; perlu periksa tarif', True),
            ('Berhasil Input', 'Error penulisan master', True),
            ('Skip - Sudah Terisi', 'Sudah terisi', True),
            ('Gagal', 'Kombinasi kode tidak ditemukan', True),
            ('Belum Diproses', '', True),
        ]
        for status, description, expected in cases:
            with self.subTest(status=status, description=description):
                self.assertEqual(app.status_input_requires_review(status, description), expected)

    def test_status_export_colors_the_entire_review_row_without_changing_data(self):
        data = pd.DataFrame([dict(ID_Transaksi=f'BANK-{i}', Kode_8_Digit=f'0101010{i}',
            Bulan='Oktober', Tahun='2026', Denda=5000, _tariff_from_kasda=i == 3,
            **{'Posting Date': '14 Oct 2026', 'Credit Transaction': 105000,
               'Nama Penghuni': 'Contoh', 'Tanggal Perjanjian Sewa': '',
               'Sewa Hunian': 100000, 'Sewa Lahan Lantai 1': 0}) for i in range(1, 7)])
        original = data.copy(deep=True)
        def detail(i, **values):
            return dict(ID_Transaksi=f'BANK-{i}', Kode_8_Digit=f'0101010{i}', Bulan='Oktober', **values)
        results = {
            'success_details': [detail(1, Sheet='LG ', Row=273, Denda=5000),
                                detail(2, Sheet='MELONG', Row=222, Denda=5000),
                                detail(3, Sheet='MELONG', Row=267, Denda=5000)],
            'skipped_details': [detail(4, Reason='Sudah terisi')],
            'failed_details': [detail(5, Reason='Error penulisan master')],
        }
        with tempfile.TemporaryDirectory() as folder:
            path = app.create_export_excel(results, data, data, data.head(0), output_dir=folder)
            wb = openpyxl.load_workbook(path)
            sheet = wb['Status_Input']
            headers = {c.value:c.column for c in sheet[1]}
            for row in range(2, 8):
                expected_review = row >= 4
                self.assertTrue(all(
                    (cell.fill.fill_type == 'solid' and cell.fill.fgColor.rgb == '00F2DCDB') == expected_review
                    for cell in sheet[row]))
                self.assertEqual(sheet.cell(row, headers['Credit Transaction']).value, 105000)
                self.assertEqual(sheet.row_dimensions[row].height, 18)
            self.assertIn('memakai Kasda', sheet.cell(4, headers['Keterangan_Input']).value)
            self.assertTrue(all(c.fill.fgColor.rgb != '00F2DCDB' for c in sheet[1]))
            self.assertTrue(all(c.fill.fgColor.rgb != '00F2DCDB' for row in wb['Kasda'] for c in row))
        pd.testing.assert_frame_equal(data, original)

    def test_verified_kasda_fees_allow_payment_but_master_identity_and_rates_stay_blank(self):
        for method in ('Referensi + tanggal + nominal',
                       'Identitas hunian + periode + tanggal + nominal'):
            with self.subTest(method=method), tempfile.TemporaryDirectory() as folder:
                root = Path(folder)
                source = root / 'master.xlsx'
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.title = 'CIGUGUR'
                for column in ('B', 'C', 'D'):
                    sheet[f'{column}3'] = '01'
                workbook.save(source)
                data = pd.DataFrame([dict(ID_Transaksi='BANK-1', Kode_8_Digit='01010101',
                    Bulan='Oktober', Tahun='2026', Nama_Kasda='Nama dari Kasda',
                    Sewa_Hunian_Kasda=300000, Sewa_Lahan_Kasda=20000,
                    Denda_Kasda=10000.25, Biaya_Lain_Kasda=5000,
                    Jumlah_Kasda=335000.25, Metode_Pencocokan=method,
                    Metode_Bayar_Kasda='QRIS',
                    **{'Posting Date': '14 Oct 2026', 'Credit Transaction': 335000.25})])
                extracted = app.extract_from_master_excel(data, {'2026': str(source)})
                eligible, manual = app.filter_unmatched_master(extracted, data.head(0))
                self.assertTrue(manual.empty)
                final = app.calculate_denda(eligible)
                row = final.iloc[0]
                self.assertEqual(row['Nama Penghuni'], '')
                self.assertEqual(row['Tanggal Perjanjian Sewa'], '')
                self.assertEqual(row['Sewa Hunian'], 300000)
                self.assertEqual(row['Sewa Lahan Lantai 1'], 20000)
                self.assertEqual(row['Denda'], 10000.25)  # Other fees are not penalties.
                valid, results = app.input_to_excel_master(final, {'2026': str(source)}, str(root / 'output'))
                self.assertEqual(results['success'], 1)
                updated = openpyxl.load_workbook(results['backup_files']['2026'])['CIGUGUR']
                for column in ('CU', 'CV', 'CW', 'CX'):
                    self.assertIsNone(updated[f'{column}3'].value)
                self.assertEqual(updated['CY3'].value, 10000.25)
                self.assertEqual(updated['DA3'].value, 'QRIS')
                self.assertIsNotNone(updated['DB3'].value)
                report = app.create_export_excel(results, valid, final, manual, output_dir=str(root / 'report'))
                status = openpyxl.load_workbook(report)['Status_Input']
                headers = {cell.value: cell.column for cell in status[1]}
                self.assertIsNone(status.cell(2, headers['Nama Penghuni']).value)
                self.assertEqual(status.cell(2, headers['Sewa Hunian']).value, 300000)
                self.assertIn('memakai Kasda', status.cell(2, headers['Keterangan_Input']).value)
                self.assertTrue(all(c.fill.fgColor.rgb == '00F2DCDB' for c in status[2]))

    def test_blank_tariff_fallback_rejects_incomplete_untrusted_or_unbalanced_kasda(self):
        base = dict(Kode_8_Digit='01010101', Bulan='Oktober', Tahun='2026',
            Nama_Kasda='Nama pada Kasda',
            Sewa_Hunian_Kasda=300000, Sewa_Lahan_Kasda=0, Denda_Kasda=10000,
            Biaya_Lain_Kasda=0, Jumlah_Kasda=310000,
            Metode_Pencocokan='Referensi + tanggal + nominal',
            **{'Nama Penghuni': '', 'Tanggal Perjanjian Sewa': '',
               'Sewa Hunian': '', 'Sewa Lahan Lantai 1': '', 'Credit Transaction': 310000})
        invalid = [
            {'Metode_Pencocokan': 'Urutan waktu + tanggal + nominal'},
            {'Metode_Pencocokan': None}, {'Denda_Kasda': None},
            {'Biaya_Lain_Kasda': 5000}, {'Jumlah_Kasda': 300000},
            {'Credit Transaction': 300000}, {'Sewa_Hunian_Kasda': 0, 'Denda_Kasda': 310000},
            {'Sewa_Lahan_Kasda': -1, 'Denda_Kasda': 10001},
            dict.fromkeys(('Nama Penghuni', 'Tanggal Perjanjian Sewa',
                           'Sewa Hunian', 'Sewa Lahan Lantai 1'), None),
        ]
        for changes in invalid:
            with self.subTest(changes=changes):
                data = pd.DataFrame([{**base, **changes}])
                eligible, manual = app.filter_unmatched_master(data, data.head(0))
                self.assertTrue(eligible.empty)
                self.assertEqual(len(manual), 1)
                self.assertFalse(manual.iloc[0]['Nama Penghuni'])

    def test_blank_tenant_name_does_not_block_payment_or_get_filled_from_bank(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / 'master.xlsx'
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = 'CIGUGUR'
            for column in ('B', 'C', 'D'):
                sheet[f'{column}3'] = '01'
            sheet['CW3'] = 300000
            workbook.save(source)
            data = pd.DataFrame([dict(ID_Transaksi='BANK-1', Kode_8_Digit='01010101',
                Bulan='Oktober', Tahun='2026', Nama_Kasda='Nama pada Kasda',
                **{'Posting Date': '14 Oct 2026', 'Credit Transaction': 310000})])
            extracted = app.extract_from_master_excel(data, {'2026': str(source)})
            eligible, manual = app.filter_unmatched_master(extracted, data.head(0))
            self.assertTrue(manual.empty)
            self.assertEqual(eligible.iloc[0]['Nama Penghuni'], '')
            _, results = app.input_to_excel_master(app.calculate_denda(eligible),
                {'2026': str(source)}, str(root / 'output'))
            self.assertEqual(results['success'], 1)
            updated = openpyxl.load_workbook(results['backup_files']['2026'])['CIGUGUR']
            self.assertIsNone(updated['CU3'].value)
            self.assertEqual(updated['CW3'].value, 300000)
            self.assertEqual(updated['CY3'].value, 10000)
            self.assertIsNotNone(updated['DB3'].value)

    def test_download_omits_technical_columns_after_matching_status_and_keeps_amounts(self):
        hidden = {
            'Urutan_Rekening_Koran', 'ID_Transaksi', 'Narasi_Clean', 'Has_SETORTUNAI',
            'Waktu_Bayar_Kasda', 'No_Referensi_Kasda', 'Nama_Kasda', 'Sheet_Kasda',
            'Tanggal_Kasda', 'Tanggal_Bayar_Kasda', 'Biaya_Lain_Kasda',
            'Nomor_Dokumen_Kasda', 'Metode_Pencocokan',
            'Jumlah_Kasda', 'Sumber_Deteksi', 'Baris_Kasda',
            'Sewa_Hunian_Kasda', 'Sewa_Lahan_Kasda', '_tariff_from_kasda',
        }
        rows = []
        for index in (1, 2):
            rows.append({**dict.fromkeys(hidden, 'data teknis'),
                'ID_Transaksi': f'BANK-{index}', 'Kode_8_Digit': '01010101',
                'Bulan': 'Oktober', 'Tahun': '2026', 'Posting Date': '14 Oct 2026',
                'Narasi': 'Baris pertama\nBaris kedua',
                'Metode_Bayar_Kasda': 'QRIS' if index == 1 else 'VA',
                'Nama Penghuni': 'Penghuni dengan nama panjang untuk pemeriksaan lebar kolom',
                'Sewa Hunian': 100000, 'Sewa Lahan Lantai 1': 0,
                'Tanggal Perjanjian Sewa': '01/01/2026', 'Denda': 5000,
                'Credit Transaction': 105000, 'Balance': 210000})
        data = pd.DataFrame(rows)
        original = data.copy(deep=True)
        results = {
            'success_details': [dict(ID_Transaksi='BANK-1', Kode_8_Digit='01010101',
                Bulan='Oktober', Sheet='CIGUGUR', Row=3, Denda=5000)],
            'skipped_details': [dict(ID_Transaksi='BANK-2', Kode_8_Digit='01010101',
                Bulan='Oktober', Reason='Sudah terisi')],
        }
        for audit in (None, data):
            with self.subTest(audit=audit is not None), tempfile.TemporaryDirectory() as folder:
                report = app.create_export_excel(results, data, data, data.head(1), audit, folder)
                workbook = openpyxl.load_workbook(report)
                self.assertEqual(workbook.sheetnames, ['Status_Input', 'Cek Manual', 'Kasda', 'Denda']
                                 + (['Pencocokan Kasda'] if audit is not None else []))
                for sheet in workbook:
                    headers = [cell.value for cell in sheet[1]]
                    self.assertFalse(hidden.intersection(headers))
                    self.assertEqual(sheet.freeze_panes, 'A2')
                    self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
                    self.assertTrue(all(dimension.height == 18 for dimension in sheet.row_dimensions.values()))
                    self.assertTrue(all(not cell.alignment.wrap_text for row in sheet for cell in row))
                status = workbook['Status_Input']
                self.assertEqual([cell.value for cell in status[1]], [
                    'No', 'Posting Date', 'Effective Date', 'Narasi', 'Debit Transaction',
                    'Credit Transaction', 'Balance', 'Kode_8_Digit', 'Rusunawa', 'Gedung',
                    'Lantai', 'No Hunian', 'Nama Penghuni', 'Tanggal Perjanjian Sewa',
                    'Sewa Hunian', 'Sewa Lahan Lantai 1', 'Denda', 'Bulan', 'Tahun',
                    'Kanal Bayar', 'Status_Input', 'Keterangan_Input',
                ])
                headers = {cell.value: cell.column for cell in status[1]}
                self.assertEqual([status.cell(row, headers['Kanal Bayar']).value for row in (2, 3)], ['QRIS', 'VA'])
                self.assertEqual([status.cell(row, headers['Status_Input']).value for row in (2, 3)],
                                 ['Berhasil Input', 'Skip - Sudah Terisi'])
                for row in (2, 3):
                    self.assertEqual(status.cell(row, headers['Credit Transaction']).value, 105000)
                    self.assertEqual(status.cell(row, headers['Credit Transaction']).number_format, '#,##0')
                    self.assertEqual(status.cell(row, headers['Kode_8_Digit']).value, '01010101')
                    self.assertEqual(status.cell(row, headers['Kode_8_Digit']).number_format, '@')
                    self.assertEqual(status.cell(row, headers['Narasi']).value, 'Baris pertama Baris kedua')
                name_column = status.cell(1, headers['Nama Penghuni']).column_letter
                self.assertGreater(status.column_dimensions[name_column].width, 18)
        pd.testing.assert_frame_equal(data, original)

    def test_blank_or_zero_tariffs_do_not_turn_the_whole_payment_into_a_penalty(self):
        data = pd.DataFrame([
            {'Kode_8_Digit': f'0101010{i}', 'Bulan': 'Oktober', 'Tahun': '2026',
             'Nama Penghuni': 'Contoh', 'Tanggal Perjanjian Sewa': '',
             'Sewa Hunian': rent, 'Sewa Lahan Lantai 1': '', 'Credit Transaction': 310000}
            for i, rent in enumerate(('', 0, 300000), 1)
        ])
        eligible, manual = app.filter_unmatched_master(data, data.head(0))
        self.assertEqual(eligible['Kode_8_Digit'].tolist(), ['01010103'])
        self.assertEqual(len(manual), 2)
        self.assertTrue(manual['Keterangan'].str.contains('Tarif sewa master kosong/nol').all())
        self.assertEqual(app.calculate_denda(eligible).iloc[0]['Denda'], 10000)

    def test_fixed_tariff_formula_cache_survives_but_changed_total_recalculates(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'master.xlsx'
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet['K3'] = '=280000+90000'
            sheet['K4'] = '=280000+90000'
            sheet['M3'] = 0
            sheet['N3'] = '=K3+M3'
            workbook.save(path)
            # Simulate the cached results written by Excel in a real input master.
            buffer = BytesIO()
            with ZipFile(path) as source, ZipFile(buffer, 'w') as out:
                for entry in source.infolist():
                    content = source.read(entry.filename)
                    if entry.filename == 'xl/worksheets/sheet1.xml':
                        root = ET.fromstring(content)
                        for cell in root.iter(NS + 'c'):
                            if cell.get('r') in ('K3', 'K4', 'N3'):
                                cell.find(NS + 'v').text = '370000'
                            if cell.get('r') in ('K3', 'K4'):
                                formula = cell.find(NS + 'f')
                                formula.set('t', 'shared')
                                formula.set('si', '0')
                                if cell.get('r')=='K3':
                                    formula.set('ref', 'K3:K4')
                                else:
                                    formula.text = None
                        content = ET.tostring(root)
                    out.writestr(entry, content)
            path.write_bytes(buffer.getvalue())
            workbook = openpyxl.load_workbook(path)
            workbook.active['M3'] = 10000
            save_master(workbook, path)
            values = openpyxl.load_workbook(path, data_only=True)
            self.assertEqual(values.active['K3'].value, 370000)
            self.assertEqual(values.active['K4'].value, 370000)
            self.assertEqual(values.active['M3'].value, 10000)
            self.assertIsNone(values.active['N3'].value)  # no stale cached 370000
            formulas = openpyxl.load_workbook(path)
            self.assertEqual(formulas.active['K3'].value, '=280000+90000')
            self.assertEqual(formulas.active['N3'].value, '=K3+M3')

    def test_all_uploaded_years_are_returned_and_repeated_input_is_skipped(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            masters = {}
            for year in ('2025', '2026'):
                path = root / f'Master_{year}.xlsx'
                workbook = openpyxl.Workbook()
                workbook.active.title = 'CIGUGUR'
                for column in ('B', 'C', 'D'):
                    workbook.active[f'{column}3'] = '01'
                workbook.save(path)
                masters[year] = str(path)
            source_contents = {y: Path(p).read_bytes() for y, p in masters.items()}
            data = pd.DataFrame([dict(ID_Transaksi='BANK-1', Kode_8_Digit='01010101',
                Bulan='Oktober', Tahun='2026', **{'Posting Date': '14 Oct 2026', 'Denda': 0})])
            _, results = app.input_to_excel_master(data.copy(), masters, str(root / 'first'))
            self.assertEqual(results['success'], 1)
            self.assertEqual(set(results['backup_files']), {'2025', '2026'})
            self.assertEqual(Path(results['backup_files']['2025']).read_bytes(), source_contents['2025'])
            for year, path in masters.items():
                self.assertEqual(Path(path).read_bytes(), source_contents[year])
            _, repeated = app.input_to_excel_master(data.copy(), results['backup_files'], str(root / 'second'))
            self.assertEqual((repeated['success'], repeated['skipped'], repeated['failed']), (0, 1, 0))
            self.assertEqual(Path(repeated['backup_files']['2026']).read_bytes(),
                             Path(results['backup_files']['2026']).read_bytes())


if __name__ == '__main__':
    unittest.main()
