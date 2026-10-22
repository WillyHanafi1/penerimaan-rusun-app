"""Read the Sewa recap and reconcile receipts with bank transactions one to one."""

from collections import Counter
from datetime import datetime, time
import re

import openpyxl
import pandas as pd


MONTHS = [
    'Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni',
    'Juli', 'Agustus', 'September', 'Oktober', 'November', 'Desember',
]
MONTH_ALIASES = {
    'jan': 1, 'januari': 1, 'feb': 2, 'februari': 2, 'mar': 3, 'maret': 3,
    'apr': 4, 'april': 4, 'mei': 5, 'may': 5, 'jun': 6, 'juni': 6,
    'jul': 7, 'juli': 7, 'agu': 8, 'ags': 8, 'agt': 8, 'agustus': 8, 'aug': 8,
    'sep': 9, 'sept': 9, 'september': 9, 'okt': 10, 'oct': 10, 'oktober': 10,
    'nov': 11, 'nop': 11, 'november': 11, 'des': 12, 'dec': 12, 'desember': 12,
}

# Nomor VA rusun: prefix 1295 + kode hunian 8 digit + periode MMYY.
# Contoh: 1295 02010417 0726 = unit 02010417, Juli 2026.
VA_REFERENCE_PATTERN = re.compile(
    r'(?<!\d)(1295(0[123]\d{6})(0[1-9]|1[0-2])(\d{2}))(?!\d)'
)


def _text(value):
    return '' if pd.isna(value) else str(value).strip()


def _key(value):
    return re.sub(r'[^A-Z0-9]', '', _text(value).upper())


def _money(value):
    try:
        number = float(_text(value).replace(',', ''))
        return round(number, 2) if pd.notna(number) else None
    except (TypeError, ValueError):
        return None


def _date(value):
    if isinstance(value, (datetime, pd.Timestamp)):
        return pd.Timestamp(value).normalize()
    value = _text(value)
    # Bank dates normalized by process_pdf use US order; named months are unambiguous.
    for fmt in ('%m/%d/%Y', '%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%d %b %Y',
                '%d%b%Y', '%d-%m-%Y', '%d/%m/%Y'):
        try:
            return pd.Timestamp(datetime.strptime(value, fmt)).normalize()
        except ValueError:
            continue
    return pd.NaT


def _period(value):
    if isinstance(value, (datetime, pd.Timestamp)):
        return MONTHS[value.month - 1], str(value.year)
    match = re.fullmatch(r'([A-Za-z]+)[\s/.-]*(\d{4}|\d{2})', _text(value))
    if match and match[1].lower() in MONTH_ALIASES:
        year = match[2] if len(match[2]) == 4 else '20' + match[2]
        return MONTHS[MONTH_ALIASES[match[1].lower()] - 1], year
    return '', ''


def _time(value):
    if pd.isna(value):
        return ''
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.strftime('%H:%M:%S')
    if isinstance(value, (float, int)) and 0 <= value < 1:
        seconds = round(value * 86400)
        return f'{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}'
    # Excel can contain native datetimes or the displayed WIB datetime copied
    # as text. Losing an anchor's time disables ordered matching for the day,
    # including otherwise valid payments after 16:00.
    for fmt in ('%H:%M:%S', '%H:%M', '%H.%M',
                '%d-%m-%Y %H:%M:%S', '%d-%m-%Y %H:%M', '%d-%m-%Y %H.%M',
                '%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M',
                '%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M',
                '%Y-%m-%dT%H:%M:%S', '%Y-%m-%dT%H:%M'):
        try:
            return datetime.strptime(_text(value), fmt).strftime('%H:%M:%S')
        except ValueError:
            continue
    return ''


def _va_reference(value):
    """Ambil nomor VA rusun yang valid dari narasi bank."""
    match = VA_REFERENCE_PATTERN.search(re.sub(r'\s+', '', _text(value)))
    return match.group(1) if match else ''


def _unit_code(rusun, gedung, lantai, unit):
    prefix = next((code for name, code in [
        ('CIGUGUR', '01'), ('CIBEUREUM', '02'), ('MELONG', '02'), ('LEUWIGAJAH', '03'),
    ] if name in rusun.upper()), '')
    building = {'A': '01', 'B': '02', 'C': '03', 'D': '04'}.get(_text(gedung).upper())
    floor = {'I': '01', 'II': '02', 'III': '03', 'IV': '04', 'V': '05'}.get(_text(lantai).upper())
    try:
        building = building or f'{int(float(gedung)):02d}'
        floor = floor or f'{int(float(lantai)):02d}'
        number = float(unit)
        if number != int(number) or not 1 <= number <= 99:
            return ''
        if prefix and building in ('01', '02', '03', '04') and floor in ('01', '02', '03', '04', '05'):
            return prefix + building + floor + f'{int(number):02d}'
    except (TypeError, ValueError):
        pass
    return ''


def _kasda_headers(values):
    """Normalize the legacy Sewa columns and the daily Inawa template."""
    keys = {_key(v): i for i, v in enumerate(values) if _text(v)}
    aliases = {
        'METODEBAYAR': ('KANALBAYAR',),
        'WAKTUBAYAR': ('WAKTUBAYARWIB',),
        'PARKIR': ('SEWALANTAIDASARPARKIR', 'SEWALAHANLANTAI1', 'SEWALAHAN'),
        'GD': ('GEDUNG',), 'LT': ('LANTAI',), 'UNIT': ('NOHUNIAN',),
    }
    for name, alternatives in aliases.items():
        if name not in keys:
            for alias in alternatives:
                if alias in keys:
                    keys[name] = keys[alias]
                    break
    # Inawa merges D:F only in the header; detail rows contain building/floor/unit.
    if 'NOMORHUNIAN' in keys and not {'GD', 'LT', 'UNIT'}.issubset(keys):
        start = keys['NOMORHUNIAN']
        keys.update(GD=start, LT=start + 1, UNIT=start + 2)
    required = {'NOURUT', 'TANGGALKASDA', 'GD', 'LT', 'UNIT', 'JUMLAH',
                'BULAN', 'WAKTUBAYAR', 'METODEBAYAR'}
    return keys if required.issubset(keys) else {}


def _daily_sheet_date(name):
    match = re.fullmatch(r'(\d{1,2})-(\d{1,2})-(\d{4})', name.strip())
    if match:
        try:
            return pd.Timestamp(datetime(int(match[3]), int(match[2]), int(match[1])))
        except ValueError:
            pass
    return None


def read_kasda_sewa(path, transaction_dates=None):
    """Read daily sheets selected by bank dates, or the legacy Sewa sheet.

    Sheet position is irrelevant. Subtotals, repeated headers, and Denda are
    excluded; receipts from other dates never enter a daily reconciliation.
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records = []
    try:
        dates = None
        if transaction_dates is not None:
            dates = {d for value in transaction_dates if pd.notna(d := _date(value))}
            if not dates:
                raise ValueError('Tanggal rekening koran tidak dapat dibaca untuk memilih sheet Kasda.')
        daily = [(name, day) for name in workbook.sheetnames if (day := _daily_sheet_date(name)) is not None]
        if daily:
            selected = [name for name, day in daily if dates is None or day in dates]
            if not selected:
                labels = ', '.join(d.strftime('%d-%m-%Y') for d in sorted(dates))
                raise ValueError(f'Sheet Kasda untuk tanggal rekening koran ({labels}) tidak tersedia.')
        else:
            selected = [s for s in workbook.sheetnames if s.strip().lower() == 'sewa']
            if not selected:
                raise ValueError("File Kasda harus memiliki sheet 'Sewa' atau sheet harian bernama D-M-YYYY.")
        for sheet_name in selected:
            headers, rusun, found_header = {}, '', False
            for row_number, values in enumerate(workbook[sheet_name].iter_rows(values_only=True), 1):
                keys = _kasda_headers(values)
                if keys:
                    headers, found_header = keys, True
                    continue
                if not values:
                    continue
                if _text(values[0]).upper().startswith('RUSUNAWA'):
                    rusun = _text(values[0])
                    headers = {}
                    continue
                if not headers:
                    continue
                def value(column):
                    index = headers.get(column)
                    return values[index] if index is not None and index < len(values) else None
                if not re.fullmatch(r'\d+(?:\.0)?', _text(value('NOURUT'))):
                    continue
                month, year = _period(value('BULAN'))
                code = _unit_code(rusun, value('GD'), value('LT'), value('UNIT'))
                kasda_date = _date(value('TANGGALKASDA'))
                # Bank posting follows the payment date, which can differ from Kasda.
                paid_date = _date(value('TANGGALSETOR')) if 'TANGGALSETOR' in headers else kasda_date
                if dates is not None and pd.notna(paid_date) and paid_date not in dates:
                    continue
                amount, method = _money(value('JUMLAH')), _key(value('METODEBAYAR'))
                issues = []
                if not code:
                    issues.append('Identitas rusun/gedung/lantai/unit tidak valid')
                if not month or not year:
                    issues.append('Periode sewa tidak valid')
                if pd.isna(paid_date):
                    issues.append('Tanggal pembayaran tidak valid')
                if amount is None or amount <= 0:
                    issues.append('Jumlah penerimaan tidak valid')
                if method not in ('QRIS', 'VA', 'MANUAL', 'TUNAI'):
                    issues.append('Metode bayar belum didukung untuk pencocokan')
                records.append({
                    'Sheet_Kasda': sheet_name, 'Baris_Kasda': row_number, 'Rusun_Kasda': rusun,
                    'Kode_8_Digit': code, 'Bulan': month, 'Tahun': year,
                    'Tanggal_Kasda': kasda_date, 'Tanggal_Bayar_Kasda': paid_date,
                    'Waktu_Bayar_Kasda': _time(value('WAKTUBAYAR')),
                    'Metode_Bayar_Kasda': method, 'No_Referensi_Kasda': _text(value('NOREFERENSI')),
                    'Nama_Kasda': _text(value('NAMAPENGHUNI')),
                    'Jumlah_Kasda': amount, 'Denda_Kasda': _money(value('DENDA')),
                    'Sewa_Hunian_Kasda': _money(value('SEWAHUNIAN')),
                    'Sewa_Lahan_Kasda': _money(value('PARKIR')),
                    'Biaya_Lain_Kasda': _money(value('BIAYALAIN')),
                    'Nomor_Dokumen_Kasda': _text(value('NOMORSKRDSTRD')),
                    'Keterangan_Pencocokan': ' | '.join(issues),
                })
            if not found_header:
                raise ValueError(f'Header penerimaan tidak ditemukan pada sheet {sheet_name}. Periksa format Kasda.')
        if not records:
            # A valid empty daily sheet is different from a missing/broken header.
            return pd.DataFrame(columns=['Sheet_Kasda', 'Baris_Kasda', 'Tanggal_Kasda',
                'Tanggal_Bayar_Kasda', 'Waktu_Bayar_Kasda', 'Metode_Bayar_Kasda',
                'No_Referensi_Kasda', 'Kode_8_Digit', 'Bulan', 'Tahun', 'Nama_Kasda',
                'Jumlah_Kasda', 'Denda_Kasda', 'Keterangan_Pencocokan'])
    finally:
        workbook.close()
    return pd.DataFrame(records)


def _ordered_pairs(bank_ids, receipt_ids, candidates, receipts):
    """Align amounts/methods in order, allowing gaps on either side.

    Only return pairs shared by every longest alignment. A repeated amount with
    several equally plausible partners stays manual, without blocking later pairs.
    """
    n, m = len(bank_ids), len(receipt_ids)
    left = [(candidates[b]['amount'], candidates[b]['method']) for b in bank_ids]
    right = [(receipts[k]['Jumlah_Kasda'], receipts[k]['Metode_Bayar_Kasda']) for k in receipt_ids]
    prefix = [[0] * (m + 1) for _ in range(n + 1)]
    suffix = [[0] * (m + 1) for _ in range(n + 1)]
    for i in range(n):
        for j in range(m):
            prefix[i + 1][j + 1] = (prefix[i][j] + 1 if left[i] == right[j]
                                     else max(prefix[i][j + 1], prefix[i + 1][j]))
    for i in range(n - 1, -1, -1):
        for j in range(m - 1, -1, -1):
            suffix[i][j] = (suffix[i + 1][j + 1] + 1 if left[i] == right[j]
                            else max(suffix[i + 1][j], suffix[i][j + 1]))
    best = prefix[n][m]
    by_position = {}
    for i in range(n):
        for j in range(m):
            if left[i] == right[j] and prefix[i][j] + 1 + suffix[i + 1][j + 1] == best:
                by_position.setdefault(prefix[i][j] + 1, []).append((bank_ids[i], receipt_ids[j]))
    return [pairs[0] for _, pairs in sorted(by_position.items()) if len(pairs) == 1]


def reconcile_kasda_sewa(detected, manual, kasda):
    """Promote matched bank rows, retaining the bank's amount, narrative and identity.

    Reference matches require equal date, payment method and amount, and uniqueness
    on both sides. Time-order matching can skip unmatched rows on either side,
    between confirmed reference anchors. Ambiguous pairs remain manual.
    """
    detected_ids = set(detected['ID_Transaksi'])
    bank = pd.concat([detected, manual], ignore_index=True).sort_values('Urutan_Rekening_Koran').reset_index(drop=True)
    audit = kasda.copy().reset_index(drop=True)
    audit['Status_Pencocokan'] = 'Belum cocok'
    audit['ID_Transaksi'] = ''
    audit['Metode_Pencocokan'] = ''
    receipts = {}
    for i, row in audit.iterrows():
        if not row['Keterangan_Pencocokan']:
            receipts[i] = row
    candidates = {}
    for i, row in bank.iterrows():
        narration = _text(row.get('Narasi')).upper()
        compact = re.sub(r'\s+', '', narration)
        va_reference = _va_reference(compact)
        # A transfer/settlement mentioning QRIS is not an individual QRIS payment.
        qris_payment = 'PAYMENTQRIS' in compact and ('RETRIBUSI' in compact or 'RUSUN' in compact)
        # Sebagian bank menampilkan pembayaran VA sebagai TRF DARI BCA/SEABANK,
        # tanpa label PAY VA. Nomor VA 16 digit tetap mengidentifikasi transaksi.
        method = 'QRIS' if qris_payment else ('VA' if 'PAYVA' in compact or va_reference else '')
        if (not method and row['ID_Transaksi'] in detected_ids
                and _text(row.get('Metode_Bayar')) in ('TUNAI', 'PEMINDAHBUKUAN')
                and not _text(row.get('_identity_issue'))
                and all(_text(row.get(c)) for c in ('Kode_8_Digit', 'Bulan', 'Tahun'))):
            method = 'MANUAL'
        credit, debit = _money(row.get('Credit Transaction')), _money(row.get('Debit Transaction'))
        date = _date(row.get('Posting Date'))
        if not method or credit is None or credit <= 0 or (debit is not None and debit > 0) or pd.isna(date):
            continue
        ref = re.search(r'(?:ESB|QRP)-[A-Z0-9]+-([A-Z0-9]{6})(?=PAYMENT|\s|$)', narration)
        if ref is None:
            ref = re.search(r'(?:ESB|QRP)-[A-Z0-9]+-([A-Z0-9]{6})(?=PAYMENT)', compact)
        if method == 'VA':
            # Channel may be TLR, MBL, ATM, etc.; the VA itself remains numeric.
            pay_va = re.search(r'PAYVA[A-Z]*(\d{10,})', compact)
            reference = va_reference or (pay_va[1] if pay_va else '')
        else:
            reference = ref[1] if ref else ''
        candidates[i] = {'date': date, 'amount': credit, 'method': method, 'ref': reference}

    edges, reference_errors, blocked_receipts = {}, {}, set()
    for b, transaction in candidates.items():
        if transaction['method'] == 'MANUAL':
            edges[b] = [k for k, receipt in receipts.items()
                        if receipt['Metode_Bayar_Kasda'] in ('MANUAL', 'TUNAI')
                        and receipt.get('Tanggal_Bayar_Kasda', receipt['Tanggal_Kasda']) == transaction['date']
                        and receipt['Jumlah_Kasda'] == transaction['amount']
                        and all(_text(receipt[c]) == _text(bank.at[b, c])
                                for c in ('Kode_8_Digit', 'Bulan', 'Tahun'))]
            continue
        if not transaction['ref']:
            continue
        edges[b], same_reference = [], []
        for k, receipt in receipts.items():
            ref = _key(receipt['No_Referensi_Kasda'])
            if transaction['method'] == 'QRIS':
                ref = ref[-6:] if len(ref) >= 6 else ''
            if (ref == transaction['ref'] and receipt.get('Tanggal_Bayar_Kasda', receipt['Tanggal_Kasda']) == transaction['date']
                    and receipt['Metode_Bayar_Kasda'] == transaction['method']):
                same_reference.append(k)
                if receipt['Jumlah_Kasda'] == transaction['amount']:
                    edges[b].append(k)
        if same_reference and not edges[b]:
            reference_errors[b] = 'Referensi cocok tetapi nominal bank dan Kasda berbeda'
            blocked_receipts.update(same_reference)
    usage = Counter(k for matches in edges.values() for k in matches)
    matches = {b: (ks[0], 'Identitas hunian + periode + tanggal + nominal'
                  if candidates[b]['method'] == 'MANUAL' else 'Referensi + tanggal + nominal') for b, ks in edges.items()
               if len(ks) == 1 and usage[ks[0]] == 1}

    for date in sorted({x['date'] for x in candidates.values()}):
        bank_day = [b for b, row in candidates.items() if row['date'] == date and row['method'] in ('QRIS', 'VA')]
        kasda_day = [k for k, row in receipts.items()
                     if row.get('Tanggal_Bayar_Kasda', row['Tanggal_Kasda']) == date
                     and row['Metode_Bayar_Kasda'] in ('QRIS', 'VA')]
        anchors = [(b, matches[b][0]) for b in bank_day if b in matches]
        if any(not receipts[k]['Waktu_Bayar_Kasda'] for _, k in anchors):
            continue
        # Conflicting chronology must not override a unique reference match.
        anchor_times = [receipts[k]['Waktu_Bayar_Kasda'] for _, k in anchors]
        if anchor_times != sorted(anchor_times):
            continue
        time_counts = Counter(receipts[k]['Waktu_Bayar_Kasda'] for k in kasda_day)
        limits = [(None, None)] + anchors + [(None, None)]
        used_receipts = {k for k, _ in matches.values()}
        for (before_bank, before_receipt), (after_bank, after_receipt) in zip(limits, limits[1:]):
            left_time = receipts[before_receipt]['Waktu_Bayar_Kasda'] if before_receipt is not None else ''
            right_time = receipts[after_receipt]['Waktu_Bayar_Kasda'] if after_receipt is not None else '24:00:00'
            bank_segment = [b for b in bank_day if b not in matches and b not in reference_errors
                            and (before_bank is None or b > before_bank)
                            and (after_bank is None or b < after_bank)]
            receipt_segment = [k for k in kasda_day if k not in used_receipts and k not in blocked_receipts
                               and time_counts[receipts[k]['Waktu_Bayar_Kasda']] == 1
                               and left_time < receipts[k]['Waktu_Bayar_Kasda'] < right_time]
            receipt_segment.sort(key=lambda k: receipts[k]['Waktu_Bayar_Kasda'])
            for b, k in _ordered_pairs(bank_segment, receipt_segment, candidates, receipts):
                matches[b] = (k, 'Urutan waktu + tanggal + nominal')
                used_receipts.add(k)

    promoted = set()
    metadata = ['Baris_Kasda', 'Waktu_Bayar_Kasda', 'No_Referensi_Kasda',
                'Nama_Kasda', 'Jumlah_Kasda', 'Denda_Kasda', 'Metode_Bayar_Kasda']
    metadata += [name for name in ['Sheet_Kasda', 'Tanggal_Kasda', 'Tanggal_Bayar_Kasda',
                 'Biaya_Lain_Kasda', 'Nomor_Dokumen_Kasda',
                 'Sewa_Hunian_Kasda', 'Sewa_Lahan_Kasda'] if name in audit]
    for column in metadata + ['Metode_Pencocokan', 'Sumber_Deteksi']:
        if column not in bank:
            bank[column] = pd.Series([None] * len(bank), dtype=object)
    for b, (k, method) in matches.items():
        receipt = receipts[k]
        bank_id = bank.at[b, 'ID_Transaksi']
        already_complete = bank_id in detected_ids and all(_text(bank.at[b, c]) for c in ['Kode_8_Digit', 'Bulan', 'Tahun'])
        if already_complete and any(_text(bank.at[b, c]) != _text(receipt[c]) for c in ['Kode_8_Digit', 'Bulan', 'Tahun']):
            audit.at[k, 'Keterangan_Pencocokan'] = 'Identitas/periode Kasda berbeda dari narasi rekening koran'
            continue
        for column in ['Kode_8_Digit', 'Bulan', 'Tahun']:
            bank.at[b, column] = receipt[column]
        for column in metadata:
            bank.at[b, column] = receipt[column]
        bank.at[b, 'Sumber_Deteksi'] = 'Narasi + Kasda Sewa' if already_complete else 'Kasda Sewa'
        bank.at[b, 'Metode_Pencocokan'] = method
        promoted.add(bank_id)
        audit.at[k, 'Status_Pencocokan'] = 'Cocok'
        audit.at[k, 'ID_Transaksi'] = bank_id
        audit.at[k, 'Metode_Pencocokan'] = method
    missing_reason = (audit['Status_Pencocokan'] != 'Cocok') & (audit['Keterangan_Pencocokan'] == '')
    audit.loc[missing_reason, 'Keterangan_Pencocokan'] = 'Tidak ditemukan pasangan urutan dengan tanggal/nominal sesuai, atau pasangan ambigu'
    for k in blocked_receipts:
        if audit.at[k, 'Status_Pencocokan'] != 'Cocok':
            audit.at[k, 'Keterangan_Pencocokan'] = 'Referensi cocok tetapi nominal bank dan Kasda berbeda'
    include = bank['ID_Transaksi'].isin(detected_ids | promoted)
    remaining = bank[~include].copy()
    if 'Keterangan' not in remaining:
        remaining['Keterangan'] = ''
    remaining['Jenis_Baris'] = 'Transaksi bank'
    remaining.loc[remaining['Keterangan'].fillna('') == '', 'Keterangan'] = 'Narasi belum terdeteksi sebagai pembayaran sewa; perlu cek manual'
    remaining.loc[remaining['ID_Transaksi'].isin([bank.at[b, 'ID_Transaksi'] for b in candidates]), 'Keterangan'] = 'Tidak ditemukan pasangan Kasda dengan tanggal/nominal sesuai, atau pasangan urutan ambigu'
    for b, reason in reference_errors.items():
        remaining.loc[remaining['ID_Transaksi'] == bank.at[b, 'ID_Transaksi'], 'Keterangan'] = reason
    unmatched_receipts = audit[audit['Status_Pencocokan'] != 'Cocok'].copy()
    if not unmatched_receipts.empty:
        unmatched_receipts['Jenis_Baris'] = 'Kasda tanpa pasangan bank'
        unmatched_receipts['Keterangan'] = unmatched_receipts['Keterangan_Pencocokan']
        # No bank credit is invented for an unmatched receipt. Its source amount
        # stays in Jumlah_Kasda so Cek Manual does not double-count bank money.
        remaining = pd.concat([remaining, unmatched_receipts], ignore_index=True)
    return bank[include].reset_index(drop=True), remaining.reset_index(drop=True), audit
